from __future__ import annotations

import asyncio
import io
from pathlib import Path
from docx import Document
from pypdf import PdfReader

from .config import EXTRACTION_TIMEOUT_SECONDS, MAX_KNOWLEDGE_FILE_BYTES, MAX_PDF_PAGES


SUPPORTED_EXTENSIONS = {".md", ".markdown", ".txt", ".pdf", ".docx"}


def is_supported_file(filename: str) -> bool:
    return Path(filename).suffix.casefold() in SUPPORTED_EXTENSIONS


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    content = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            content.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(item for item in content if item.strip()).strip()


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f"PDF documents are limited to {MAX_PDF_PAGES} pages.")
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


async def extract_file(filename: str, data: bytes) -> str:
    if len(data) > MAX_KNOWLEDGE_FILE_BYTES:
        raise ValueError("The uploaded document must be smaller than 5 MB.")
    extension = Path(filename).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Upload a Markdown, text, PDF, or Word (.docx) document.")
    if extension in {".md", ".markdown", ".txt"}:
        try:
            return data.decode("utf-8-sig").strip()
        except UnicodeDecodeError as error:
            raise ValueError("The text document must use UTF-8 encoding.") from error
    extractor = _extract_docx if extension == ".docx" else _extract_pdf
    label = "Word document" if extension == ".docx" else "PDF document"
    try:
        return await asyncio.wait_for(
            asyncio.to_thread(extractor, data), timeout=EXTRACTION_TIMEOUT_SECONDS
        )
    except asyncio.TimeoutError as error:
        raise ValueError(f"{label} extraction timed out.") from error
